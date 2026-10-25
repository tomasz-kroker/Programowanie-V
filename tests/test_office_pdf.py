import unittest
import tempfile
from pathlib import Path
from docx import Document
from unittest.mock import patch

from converters.office_pdf import docx_to_pdf, odt_to_pdf


class TestOfficePdf(unittest.TestCase):

    def test_docx_to_pdf(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "test.docx"
            destination = Path(folder) / "test.pdf"

            document = Document()
            document.add_paragraph("Zażółć gęślą jaźń")
            document.save(source)

            docx_to_pdf(source, destination)

            self.assertTrue(destination.is_file())
            self.assertGreater(destination.stat().st_size, 0)

            with destination.open("rb") as file:
                self.assertEqual(file.read(5), b"%PDF-")

    def test_odt_to_pdf(self):
        project_dir = Path(__file__).resolve().parent.parent
        source = project_dir / "sample_files" / "przyklad.odt"

        with tempfile.TemporaryDirectory() as folder:
            destination = Path(folder) / "wynik.pdf"

            odt_to_pdf(source, destination)

            self.assertTrue(destination.is_file())
            self.assertGreater(destination.stat().st_size, 0)

            with destination.open("rb") as file:
                self.assertEqual(file.read(5), b"%PDF-")

    def test_existing_pdf(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "test.docx"
            destination = Path(folder) / "test.pdf"

            document = Document()
            document.add_paragraph("Test")
            document.save(source)

            destination.write_bytes(b"Istniejacy plik")

            with self.assertRaises(FileExistsError):
                docx_to_pdf(source, destination)

            self.assertEqual(
                destination.read_bytes(),
                b"Istniejacy plik"
            )

    def test_invalid_extension(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "test.txt"
            destination = Path(folder) / "test.pdf"

            source.write_text("Test", encoding="utf-8")

            with self.assertRaises(ValueError):
                docx_to_pdf(source, destination)


    
    def test_missing_libreoffice(self):

        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "test.docx"
            destination = Path(folder) / "test.pdf"

            document = Document()
            document.add_paragraph("Test")
            document.save(source)

            with patch(
                "converters.office_pdf.find_libreoffice",
                return_value=None
            ):
                with self.assertRaises(FileNotFoundError):
                    docx_to_pdf(source, destination)

            self.assertFalse(destination.exists())



if __name__ == "__main__":
    unittest.main()
