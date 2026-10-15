import unittest
import tempfile
from pathlib import Path
from docx import Document

from converters.txt_docx import txt_to_docx, docx_to_txt


class TestTxtDocx(unittest.TestCase):

    def test_txt_to_docx(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "test.txt"
            destination = Path(folder) / "test.docx"

            source.write_text(
                "Pierwszy akapit\nDrugi akapit",
                encoding="utf-8"
            )

            txt_to_docx(source, destination)

            document = Document(destination)

            self.assertEqual(
                document.paragraphs[0].text,
                "Pierwszy akapit"
            )
            self.assertEqual(
                document.paragraphs[1].text,
                "Drugi akapit"
            )

    def test_docx_to_txt(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "test.docx"
            destination = Path(folder) / "test.txt"

            document = Document()
            document.add_paragraph("Akapit pierwszy")
            document.add_paragraph("Akapit drugi")
            document.save(source)

            docx_to_txt(source, destination)

            content = destination.read_text(encoding="utf-8")

            self.assertEqual(
                content,
                "Akapit pierwszy\nAkapit drugi"
            )

    def test_polish_characters_and_paragraphs(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "polskie.txt"
            docx_file = Path(folder) / "polskie.docx"
            output = Path(folder) / "wynik.txt"

            original = "Łódź, Żory, Kraków\n\nZażółć gęślą jaźń"

            source.write_text(original, encoding="utf-8")

            txt_to_docx(source, docx_file)
            docx_to_txt(docx_file, output)

            result = output.read_text(encoding="utf-8")

            self.assertEqual(result, original)


if __name__ == "__main__":
    unittest.main()
