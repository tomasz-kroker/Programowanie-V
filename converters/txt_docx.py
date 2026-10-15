from pathlib import Path
from docx import Document


def txt_to_docx(source: Path, destination: Path) -> Path:

    with source.open("r", encoding="utf-8-sig") as file:
        content = file.read()

    document = Document()

    for line in content.split("\n"):
        document.add_paragraph(line.rstrip("\r"))

    document.save(destination)

    return destination


def docx_to_txt(source: Path, destination: Path) -> Path:

    document = Document(source)

    paragraphs = []

    for paragraph in document.paragraphs:
        paragraphs.append(paragraph.text)

    content = "\n".join(paragraphs)

    with destination.open("w", encoding="utf-8") as file:
        file.write(content)

    return destination

